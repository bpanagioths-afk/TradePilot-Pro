import csv
import io
from datetime import datetime

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.deps import get_current_user, get_db
from app.models.trade import Trade
from app.models.user import User
from app.reports.pdf_report import build_trades_pdf_report


router = APIRouter(
    prefix="/exports",
    tags=["Exports"],
)


EXPORT_COLUMNS = [
    ("ID", "id"),
    ("Symbol", "symbol"),
    ("Direction", "direction"),
    ("Entry Price", "entry_price"),
    ("Stop Loss", "stop_loss"),
    ("Take Profit", "take_profit"),
    ("Exit Price", "exit_price"),
    ("Lot Size", "lot_size"),
    ("Profit Money", "profit_money"),
    ("Profit Pips", "profit_pips"),
    ("Movement Value", "movement_value"),
    ("Movement Unit", "movement_unit"),
    ("Asset Class", "asset_class"),
    ("Risk Reward", "risk_reward"),
    ("Duration Minutes", "duration_minutes"),
    ("Result", "is_win"),
    ("Open Time", "open_time"),
    ("Close Time", "close_time"),
    ("Session", "session_name"),
    ("Trading System ID", "trading_system_id"),
    ("Psychology State ID", "psychology_state_id"),
    ("MT5 Account ID", "mt5_account_id"),
    ("MT5 Ticket", "mt5_ticket"),
    ("MT5 Position ID", "mt5_position_id"),
    ("Imported From MT5", "imported_from_mt5"),
    ("TradingView Link", "tradingview_link"),
    ("Screenshot Path", "screenshot_path"),
    ("Notes", "notes"),
]


def _get_user_trades(
    db: Session,
    user_id: int,
) -> list[Trade]:
    return (
        db.query(Trade)
        .filter(
            Trade.user_id == user_id,
        )
        .order_by(
            Trade.open_time.desc(),
            Trade.id.desc(),
        )
        .all()
    )


def _format_export_value(value):
    if isinstance(value, datetime):
        return value.isoformat(
            sep=" ",
            timespec="seconds",
        )

    if isinstance(value, bool):
        return "Yes" if value else "No"

    if value is None:
        return ""

    return value


def _trade_export_row(
    trade: Trade,
) -> list:
    values = []

    for _, attribute_name in EXPORT_COLUMNS:
        value = getattr(
            trade,
            attribute_name,
            None,
        )

        if attribute_name == "is_win":
            if value == 1:
                value = "WIN"
            elif value == 0:
                value = "LOSS"
            else:
                value = ""

        values.append(
            _format_export_value(value)
        )

    return values


@router.get("/trades/csv")
def export_trades_csv(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    trades = _get_user_trades(
        db,
        current_user.id,
    )

    text_buffer = io.StringIO(
        newline=""
    )

    writer = csv.writer(
        text_buffer
    )

    writer.writerow(
        [
            column_title
            for column_title, _ in EXPORT_COLUMNS
        ]
    )

    for trade in trades:
        writer.writerow(
            _trade_export_row(trade)
        )

    csv_bytes = (
        "\ufeff"
        + text_buffer.getvalue()
    ).encode("utf-8")

    output = io.BytesIO(csv_bytes)

    filename = (
        "tradepilot_trades_"
        f"{datetime.now():%Y%m%d_%H%M%S}.csv"
    )

    return StreamingResponse(
        output,
        media_type=(
            "text/csv; charset=utf-8"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


@router.get("/trades/pdf")
def export_trades_pdf(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    trades = _get_user_trades(
        db,
        current_user.id,
    )

    pdf_buffer = (
        build_trades_pdf_report(
            trades
        )
    )

    filename = (
        "tradepilot_report_"
        f"{datetime.now():%Y%m%d_%H%M%S}.pdf"
    )

    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


@router.get("/trades/excel")
def export_trades_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    trades = _get_user_trades(
        db,
        current_user.id,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Trades"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="0F172A",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    headers = [
        column_title
        for column_title, _ in EXPORT_COLUMNS
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for trade in trades:
        worksheet.append(
            _trade_export_row(trade)
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        maximum_length = len(header)

        for cell in worksheet[
            get_column_letter(column_index)
        ]:
            cell_value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            maximum_length = max(
                maximum_length,
                len(cell_value),
            )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(
            maximum_length + 3,
            45,
        )

    output = io.BytesIO()

    workbook.save(output)
    output.seek(0)

    filename = (
        "tradepilot_trades_"
        f"{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )